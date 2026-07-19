#!/usr/bin/env python3
"""Gera SQL de importacao de historico de manutencao a partir da planilha
'Frota_ZANATTA_VDH_vXX.xlsx' (aba 'Controle_Manutenção').

Cada item da coluna "Detalhamento Despesa" separado por "/" vira uma linha
em work_order_items (a planilha nao tem custo por item, so o total da OS).

Uso:
  python3 scripts/import-maintenance-xlsx.py <planilha.xlsx> <vehicles.json> > db/seed/import.sql

<vehicles.json> e o resultado de:
  npx wrangler d1 execute sigf-db --local --json \
    --command="SELECT id, plate, asset_code FROM vehicles" > vehicles.json
"""

import json
import re
import sys
from datetime import datetime

import openpyxl

SHEET_NAME = "Controle_Manutenção"
NCOLS = 22

TIPO_MAP = {
    "PREVENTIVA": "preventiva",
    "CORRETIVA": "corretiva",
    "CORRERTIVA": "corretiva",
    "CORRETIVO": "corretiva",
    "CORRETICA": "corretiva",
    "CORRETICVA": "corretiva",
    "PERIODICA": "periodica",
    "PERIODICO": "periodica",
    "PNEUS": "pneus",
    "RECALL": "recall",
    "ACESSÓRIOS": "acessorios",
    "RECUPERAÇÃO": "outro",
}

# Códigos que sabidamente não correspondem a nenhum veículo real (nem ativo
# nem descontinuado) — nunca tentar adivinhar, so registrar como ignorado.
KNOWN_UNRESOLVABLE_CODES = {"CAMBIO"}


def normalize_plate(raw):
    return re.sub(r"[^A-Z0-9]", "", str(raw).upper())


def normalize_code(raw):
    s = str(raw).strip().upper()
    if s.endswith(".0"):
        s = s[:-2]
    return s


def sql_str(value):
    if value is None:
        return "NULL"
    return "'" + str(value).replace("'", "''") + "'"


def sql_int(value):
    if value is None:
        return "NULL"
    return str(int(value))


def clean_text(v):
    if v is None:
        return None
    s = str(v).strip()
    if not s or s in ("-", "--", "---", "#N/A"):
        return None
    return s


def parse_valor_cents(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return round(v * 100)
    s = str(v).strip()
    if not s or s in ("-", "--", "---"):
        return None
    s = s.replace(".", "").replace(",", ".") if "," in s else s
    try:
        return round(float(s) * 100)
    except ValueError:
        return None


def parse_os_number(v):
    if v is None:
        return None
    if isinstance(v, float):
        v = int(v)
    s = str(v).strip()
    if not s or s in ("-", "--", "---"):
        return None
    return s


def parse_invoice_number(v):
    if v is None:
        return None
    s = str(v).strip()
    if not s or s in ("-", "--", "---"):
        return None
    return " / ".join(part.strip() for part in s.splitlines() if part.strip())


def parse_odometer(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return round(v)
    s = str(v).strip()
    if not s or s in ("-", "--", "---"):
        return None
    try:
        return round(float(s))
    except ValueError:
        return None


LADO_SLASH_RE = re.compile(r"(?<=\b[LDERlder])/(?=[LDERlder]\b)")
CODE_SLASH_RE = re.compile(r"(?<=\d)/(?=\d?[A-Za-z]\b)")
SLASH_PLACEHOLDER = ""


def split_items(text):
    """Divide por "/" tratando cada pedaco como um item, mas preservando
    abreviacoes de lado (L/D, L/E, D/E) e codigos tipo lampada (P21/5W),
    que nao sao quebras de item de verdade."""
    protected = LADO_SLASH_RE.sub(SLASH_PLACEHOLDER, text)
    protected = CODE_SLASH_RE.sub(SLASH_PLACEHOLDER, protected)
    parts = [p.strip().replace(SLASH_PLACEHOLDER, "/") for p in protected.split("/")]
    return [p for p in parts if p]


def parse_tipo(v):
    if v is None:
        return None
    s = str(v).strip().upper()
    return TIPO_MAP.get(s)


def build_vehicle_lookup(vehicles):
    by_plate = {}
    by_code = {}
    for v in vehicles:
        if v.get("plate"):
            by_plate[normalize_plate(v["plate"])] = v["id"]
        if v.get("asset_code"):
            code = normalize_code(v["asset_code"])
            by_code[code] = v["id"]
            by_code[code.zfill(4)] = v["id"]
            by_code.setdefault(code.lstrip("0") or "0", v["id"])
    return by_plate, by_code


def resolve_vehicle(placa_raw, codigo_raw, by_plate, by_code):
    if placa_raw and str(placa_raw).strip().upper() not in ("#N/A", ""):
        vid = by_plate.get(normalize_plate(placa_raw))
        if vid:
            return vid
    if codigo_raw not in (None, "#N/A", ""):
        code_norm = normalize_code(codigo_raw)
        vid = by_code.get(code_norm) or by_code.get(code_norm.zfill(4)) or by_code.get(
            code_norm.lstrip("0") or "0"
        )
        if vid:
            return vid
    return None


def main():
    if len(sys.argv) != 3:
        print(__doc__, file=sys.stderr)
        sys.exit(1)

    xlsx_path, vehicles_json_path = sys.argv[1], sys.argv[2]

    with open(vehicles_json_path, encoding="utf-8") as f:
        vjson = json.load(f)
    vehicles = vjson[0]["results"]
    by_plate, by_code = build_vehicle_lookup(vehicles)

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[SHEET_NAME]

    statements = [
        "-- Gerado automaticamente por scripts/import-maintenance-xlsx.py",
        f"-- Fonte: {xlsx_path}",
        f"-- Gerado em: {datetime.now().isoformat(timespec='seconds')}",
        "",
    ]

    imported = 0
    skipped_blank = 0
    skipped_unresolved = []
    dedup_seen = set()
    dupes_skipped = 0

    # IDs explicitos (nao depende de last_insert_rowid(), que muda a cada
    # linha inserida — inclusive nas linhas de work_order_items — e por
    # isso nao pode ser reaproveitado com seguranca dentro de um INSERT
    # multi-linha). Comeca de um offset alto para nunca colidir com OS
    # criadas manualmente no sistema (que comecam do id 1).
    next_id = 100001

    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(1, NCOLS + 1)]
        # 0 ident pneus, 1 data, 2 codigo, 3 placa, 4 modelo, 5 anofab, 6 anomod,
        # 7 km veic, 8 km prox, 9 motorista, 10 solicitado, 11 aprovado, 12 nos,
        # 13 nf, 14 valor, 15 forma pagto, 16 prazo pagto, 17 realizado por,
        # 18 tipo, 19 (vazia), 20 obra, 21 detalhamento
        data, codigo, placa = vals[1], vals[2], vals[3]
        valor, detalhamento = vals[14], vals[21]

        if all(x is None for x in (data, codigo, valor, detalhamento)):
            skipped_blank += 1
            continue

        codigo_norm = normalize_code(codigo) if codigo is not None else None
        if codigo_norm in KNOWN_UNRESOLVABLE_CODES:
            skipped_unresolved.append((r, placa, codigo, detalhamento))
            continue

        vehicle_id = resolve_vehicle(placa, codigo, by_plate, by_code)
        if vehicle_id is None:
            skipped_unresolved.append((r, placa, codigo, detalhamento))
            continue

        opened_at = None
        if isinstance(data, datetime):
            opened_at = data.strftime("%Y-%m-%d")

        valor_cents = parse_valor_cents(valor)
        detalhe_clean = clean_text(detalhamento)

        dedup_key = (vehicle_id, opened_at, valor_cents, detalhe_clean)
        if dedup_key in dedup_seen:
            dupes_skipped += 1
            continue
        dedup_seen.add(dedup_key)

        problem_description = detalhe_clean or "Manutencao (sem detalhamento na planilha)"
        maintenance_type = parse_tipo(vals[18])
        workshop = clean_text(vals[17])
        requested_by = clean_text(vals[10])
        approved_by = clean_text(vals[11])
        payment_method = clean_text(vals[15])
        payment_term = clean_text(vals[16])
        project_client = clean_text(vals[20])
        os_number = parse_os_number(vals[12])
        invoice_number = parse_invoice_number(vals[13])
        odometer_at_service = parse_odometer(vals[7])

        notes_parts = []
        ident_pneus = clean_text(vals[0])
        if ident_pneus:
            notes_parts.append(f"Identificacao pneus: {ident_pneus}")
        notes_parts.append(
            "Importado do historico de manutencao (Frota_ZANATTA_VDH, aba Controle_Manutencao)."
        )
        notes = " | ".join(notes_parts)

        wo_id = next_id
        next_id += 1

        wo_values = {
            "id": wo_id,
            "vehicle_id": vehicle_id,
            "status": "encerrada",
            "problem_description": problem_description,
            "workshop": workshop,
            "requested_by": requested_by,
            "approved_by": approved_by,
            "payment_method": payment_method,
            "final_cost_cents": valor_cents,
            "opened_at": opened_at,
            "closed_at": opened_at,
            "notes": notes,
            "maintenance_type": maintenance_type,
            "os_number": os_number,
            "invoice_number": invoice_number,
            "payment_term": payment_term,
            "project_client": project_client,
            "odometer_at_service": odometer_at_service,
        }

        int_fields = {"id", "vehicle_id", "final_cost_cents", "odometer_at_service"}
        columns = ", ".join(wo_values.keys())
        formatted = []
        for key, val in wo_values.items():
            if key in int_fields:
                formatted.append(sql_int(val))
            elif key == "opened_at":
                formatted.append(f"COALESCE({sql_str(val)}, datetime('now'))")
            else:
                formatted.append(sql_str(val))

        statements.append(
            f"INSERT INTO work_orders ({columns}) VALUES ({', '.join(formatted)});"
        )

        if detalhe_clean:
            items = split_items(detalhe_clean)
            if items:
                values_rows = ", ".join(f"({wo_id}, {sql_str(item)}, 1, 0)" for item in items)
                statements.append(
                    "INSERT INTO work_order_items (work_order_id, description, quantity, unit_cost_cents) "
                    f"VALUES {values_rows};"
                )

        imported += 1

    statements.append("")
    statements.append(f"-- Total importado: {imported}")
    statements.append(f"-- Linhas em branco/formula quebrada ignoradas: {skipped_blank}")
    statements.append(f"-- Duplicatas exatas ignoradas: {dupes_skipped}")
    statements.append(f"-- Nao identificados (nao importados): {len(skipped_unresolved)}")

    print("\n".join(statements))

    if skipped_unresolved:
        print(f"\n-- Registros NAO importados (sem veiculo correspondente):", file=sys.stderr)
        for r, placa, codigo, detalhe in skipped_unresolved:
            print(f"--   linha {r}: placa={placa!r} codigo={codigo!r} detalhe={detalhe!r}", file=sys.stderr)

    print(f"\nResumo: {imported} importados, {skipped_blank} em branco, "
          f"{dupes_skipped} duplicatas, {len(skipped_unresolved)} nao identificados",
          file=sys.stderr)


if __name__ == "__main__":
    main()
