#!/usr/bin/env python3
"""Gera SQL de importacao da tabela vehicles a partir da planilha
'SIGF_Cadastro_Veiculos.xlsx' (aba 'Cadastro de veiculos').

Uso:
  python3 scripts/import-vehicles-xlsx.py <planilha.xlsx> > db/seed/import.sql
"""

import re
import sys
from datetime import datetime

import openpyxl

NAO_INFORMADO = {"não informado", "nao informado", ""}

HEADER_MAP = {
    "Placa / Identificador": "identifier",
    "Nome / Modelo": "model",
    "Tipo": "vehicle_type",
    "Responsável": "responsible",
    "Centro de custo": "cost_center",
    "Hodômetro": "odometer",
    "Referência hodômetro": "odometer_reference_date",
    "Status": "status_raw",
    "Código patrimonial": "asset_code",
    "Tem placa": "has_plate",
    "Placa cadastro": "registered_plate",
    "UF / Base": "uf_base",
    "Ano fabricação": "manufacture_year",
    "Ano modelo": "model_year",
    "RENAVAM": "renavam",
    "Chassi": "chassis",
    "Proprietário": "owner_name",
    "Categoria": "fleet_class",
    "Apelido": "nickname",
    "Lote importação": "import_batch",
    "Arquivo origem": "source_file",
}


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        v = value.strip()
        if v.lower() in NAO_INFORMADO:
            return None
        return v
    return value


def clean_int(value):
    v = clean(value)
    if v is None:
        return None
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def normalize_plate(raw):
    upper = raw.upper()
    return re.sub(r"[^A-Z0-9]", "", upper)


def sql_str(value):
    if value is None:
        return "NULL"
    return "'" + str(value).replace("'", "''") + "'"


def sql_int(value):
    if value is None:
        return "NULL"
    return str(int(value))


def derive_category(row):
    if clean(row.get("cost_center")) == "PARTICULAR":
        return "particular"
    if clean(row.get("vehicle_type")) == "Equipamento":
        return "equipamento"
    return "veiculo"


def main():
    if len(sys.argv) != 2:
        print(__doc__, file=sys.stderr)
        sys.exit(1)

    path = sys.argv[1]
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Cadastro de veículos"]

    headers = [ws.cell(row=4, column=c).value for c in range(1, ws.max_column + 1)]
    col_index = {h: i for i, h in enumerate(headers) if h in HEADER_MAP}

    statements = [
        "-- Gerado automaticamente por scripts/import-vehicles-xlsx.py",
        f"-- Fonte: {path}",
        f"-- Gerado em: {datetime.now().isoformat(timespec='seconds')}",
        "",
    ]

    total = 0
    for r in range(5, ws.max_row + 1):
        raw = {
            key: ws.cell(row=r, column=col_index[header] + 1).value
            for header, key in HEADER_MAP.items()
            if header in col_index
        }
        if all(v is None for v in raw.values()):
            continue

        identifier = clean(raw.get("identifier"))
        if not identifier:
            continue

        has_plate = clean(raw.get("has_plate")) == "Sim"
        plate = normalize_plate(identifier) if has_plate else None
        asset_code = clean(raw.get("asset_code")) or (None if has_plate else identifier)

        odometer_ref = raw.get("odometer_reference_date")
        odometer_ref_str = None
        if isinstance(odometer_ref, datetime):
            odometer_ref_str = odometer_ref.strftime("%Y-%m-%d")

        category = derive_category(raw)
        odometer = clean_int(raw.get("odometer")) or 0

        notes_parts = []
        batch = clean(raw.get("import_batch"))
        source = clean(raw.get("source_file"))
        if batch or source:
            notes_parts.append(
                f"Importado de {source or 'planilha'}"
                + (f" (lote {batch})" if batch else "")
            )

        values = {
            "category": category,
            "plate": plate,
            "registered_plate": clean(raw.get("registered_plate")),
            "model": clean(raw.get("model")) or identifier,
            "vehicle_type": clean(raw.get("vehicle_type")),
            "nickname": clean(raw.get("nickname")),
            "responsible": clean(raw.get("responsible")),
            "cost_center": clean(raw.get("cost_center")),
            "status": "disponivel",
            "asset_code": asset_code,
            "renavam": clean(raw.get("renavam")),
            "chassis": clean(raw.get("chassis")),
            "manufacture_year": clean_int(raw.get("manufacture_year")),
            "model_year": clean_int(raw.get("model_year")),
            "owner_name": clean(raw.get("owner_name")),
            "odometer": odometer,
            "odometer_reference_date": odometer_ref_str,
            "uf_base": clean(raw.get("uf_base")),
            "fleet_class": clean(raw.get("fleet_class")),
            "notes": " | ".join(notes_parts) or None,
        }

        columns = ", ".join(values.keys())
        formatted = []
        for key, val in values.items():
            if key in ("manufacture_year", "model_year", "odometer"):
                formatted.append(sql_int(val))
            else:
                formatted.append(sql_str(val))

        statements.append(
            f"INSERT INTO vehicles ({columns}) VALUES ({', '.join(formatted)});"
        )
        total += 1

    statements.append("")
    statements.append(f"-- Total de registros: {total}")

    print("\n".join(statements))


if __name__ == "__main__":
    main()
